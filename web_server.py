import os
import sys
import re
import csv
import json
import time
import io
import zipfile
import threading
import hashlib
from datetime import datetime
from typing import List, Dict, Any, Optional

from flask import Flask, jsonify, request, Response, send_file, render_template, session
from imap_tools import MailBox
import pymupdf
import docx

from services.config_manager import ConfigManager
from services.email_watcher import EmailWatcher
from services.folder_watcher import FolderWatcher
from services.notification_service import send_windows_toast, send_telegram_alert
from services.startup_manager import is_startup_enabled, enable_startup, disable_startup

app_dir = os.path.dirname(os.path.abspath(__file__))
config_file = os.path.join(app_dir, "config.json")
config_manager = ConfigManager(config_file)

app = Flask(__name__, template_folder=os.path.join(app_dir, "templates"))
app.secret_key = os.environ.get("RESUMEFLOW_SECRET_KEY", "rf_enterprise_master_key_parvez_7221")
MASTER_PASSWORD = os.environ.get("RESUMEFLOW_MASTER_PASSWORD", "Parvez7221")

log_messages = []
log_lock = threading.Lock()
active_watcher: Optional[Any] = None
email_watcher: Optional[Any] = None
watcher_lock = threading.Lock()

SAFE_DOC_EXTENSIONS = ('.pdf', '.docx', '.doc', '.rtf', '.txt')

MANIFEST_HEADERS = [
    "Downloaded At", "Company", "Applicant Name", "Applicant Email",
    "Email Subject", "Filename", "File Size (KB)", "File Path",
    "SHA256", "Status", "Notes"
]

@app.before_request
def require_master_password():
    # Exclude non-API routes (root UI, static files)
    if not request.path.startswith("/api/"):
        return None
    # Exclude public auth endpoints
    if request.path in ["/api/login", "/api/auth_status"]:
        return None
    # Enforce session authentication
    if not session.get("authenticated"):
        return jsonify({
            "success": False,
            "auth_required": True,
            "error": "Authentication required. Please enter master password."
        }), 401
    return None

@app.route("/api/login", methods=["POST"])
def api_login():
    data = request.json or {}
    pwd = (data.get("password") or "").strip()
    if pwd == MASTER_PASSWORD:
        session["authenticated"] = True
        add_log("Master security key verified. Dashboard unlocked.", "SUCCESS")
        return jsonify({"success": True, "message": "Authenticated successfully."})
    else:
        add_log("Failed dashboard unlock attempt with invalid security key.", "WARNING")
        return jsonify({"success": False, "error": "Invalid master password. Access denied."}), 401

@app.route("/api/logout", methods=["POST"])
def api_logout():
    session.pop("authenticated", None)
    add_log("Dashboard locked / session logged out.", "INFO")
    return jsonify({"success": True, "message": "Logged out successfully."})

@app.route("/api/auth_status", methods=["GET"])
def api_auth_status():
    is_auth = bool(session.get("authenticated", False))
    return jsonify({"authenticated": is_auth})

def add_log(message: str, level: str = "INFO"):
    timestamp = datetime.now().strftime("%H:%M:%S")
    entry = {"timestamp": timestamp, "level": level, "message": message}
    with log_lock:
        log_messages.append(entry)
        if len(log_messages) > 300:
            log_messages.pop(0)

def get_manifest_records(base_dir: str) -> List[Dict[str, Any]]:
    records = []
    manifest_path = os.path.join(base_dir, "resumes_manifest.csv")
    if os.path.exists(manifest_path):
        try:
            with open(manifest_path, "r", encoding="utf-8") as f:
                reader = csv.DictReader(f)
                for row in reader:
                    if not row.get("Status"):
                        row["Status"] = "New"
                    if not row.get("Notes"):
                        row["Notes"] = ""
                    records.append(row)
        except Exception as e:
            add_log(f"Notice reading manifest: {e}", "WARNING")
    return list(reversed(records))

def update_manifest_row(base_dir: str, file_path: str, new_status: str, new_notes: str) -> bool:
    manifest_path = os.path.join(base_dir, "resumes_manifest.csv")
    if not os.path.exists(manifest_path):
        return False
    rows = []
    updated = False
    try:
        with open(manifest_path, "r", encoding="utf-8") as f:
            reader = csv.reader(f)
            header = next(reader, None)
            if not header:
                return False
            col_map = {col.strip(): i for i, col in enumerate(header)}
            path_idx = col_map.get("File Path", 7)
            status_idx = col_map.get("Status", -1)
            notes_idx = col_map.get("Notes", -1)

            if status_idx == -1:
                header.append("Status")
                status_idx = len(header) - 1
            if notes_idx == -1:
                header.append("Notes")
                notes_idx = len(header) - 1

            for row in reader:
                if not row:
                    continue
                while len(row) < len(header):
                    row.append("")
                row_path = row[path_idx] if path_idx < len(row) else ""
                
                norm_target = os.path.normcase(os.path.abspath(file_path))
                norm_row = os.path.normcase(os.path.abspath(row_path)) if row_path else ""
                if norm_target == norm_row or os.path.basename(file_path) == os.path.basename(row_path):
                    row[status_idx] = new_status
                    row[notes_idx] = new_notes
                    updated = True
                rows.append(row)

        if updated:
            with open(manifest_path, "w", newline="", encoding="utf-8") as f:
                writer = csv.writer(f)
                writer.writerow(header)
                writer.writerows(rows)
        return updated
    except Exception as e:
        add_log(f"Error updating candidate status: {e}", "ERROR")
        return False

# Document reader & writer engine
def read_resume_document(file_path: str):
    if not os.path.exists(file_path):
        return {"success": False, "error": "File does not exist."}
    
    ext = os.path.splitext(file_path)[1].lower()
    text = ""
    try:
        if ext == ".docx":
            doc = docx.Document(file_path)
            paragraphs = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                    if row_text:
                        paragraphs.append(row_text)
            text = "\n".join(paragraphs)
        elif ext == ".pdf":
            doc = pymupdf.open(file_path)
            pages_text = []
            for page in doc:
                pages_text.append(page.get_text())
            doc.close()
            text = "\n".join(pages_text).strip()
        elif ext in [".txt", ".rtf"]:
            try:
                with open(file_path, "r", encoding="utf-8") as f:
                    text = f.read()
            except UnicodeDecodeError:
                with open(file_path, "r", encoding="cp1252", errors="ignore") as f:
                    text = f.read()
        else:
            return {"success": False, "error": f"Unsupported format '{ext}' for editing."}

        size_kb = round(os.path.getsize(file_path) / 1024, 1)
        return {
            "success": True,
            "filename": os.path.basename(file_path),
            "file_path": file_path,
            "format": ext.replace(".", "").upper(),
            "size_kb": size_kb,
            "content": text
        }
    except Exception as e:
        return {"success": False, "error": f"Could not read document: {str(e)}"}

def save_resume_document(file_path: str, new_content: str, save_as_new: bool = False):
    if not os.path.exists(file_path) and not save_as_new:
        return {"success": False, "error": "Original file not found."}

    ext = os.path.splitext(file_path)[1].lower()
    target_path = file_path

    if save_as_new:
        base, extension = os.path.splitext(file_path)
        counter = 1
        target_path = f"{base}_edited{extension}"
        while os.path.exists(target_path):
            target_path = f"{base}_edited_{counter}{extension}"
            counter += 1

    try:
        if ext == ".docx":
            new_doc = docx.Document()
            for line in new_content.split("\n"):
                new_doc.add_paragraph(line)
            new_doc.save(target_path)
        elif ext == ".pdf":
            doc = pymupdf.open()
            margin = 50
            page_width = 595
            page_height = 842
            lines = new_content.split("\n")
            current_page = doc.new_page(width=page_width, height=page_height)
            y = margin + 15
            line_height = 14

            for line in lines:
                while len(line) > 95:
                    chunk = line[:95]
                    line = line[95:]
                    if y > page_height - margin:
                        current_page = doc.new_page(width=page_width, height=page_height)
                        y = margin + 15
                    current_page.insert_text((margin, y), chunk, fontsize=10.5, fontname="helv")
                    y += line_height

                if y > page_height - margin:
                    current_page = doc.new_page(width=page_width, height=page_height)
                    y = margin + 15
                current_page.insert_text((margin, y), line, fontsize=10.5, fontname="helv")
                y += line_height

            doc.save(target_path)
            doc.close()
        else:
            with open(target_path, "w", encoding="utf-8") as f:
                f.write(new_content)

        size_kb = round(os.path.getsize(target_path) / 1024, 1)
        with open(target_path, "rb") as f:
            new_sha = hashlib.sha256(f.read()).hexdigest()

        return {
            "success": True,
            "file_path": target_path,
            "filename": os.path.basename(target_path),
            "size_kb": size_kb,
            "sha256": new_sha,
            "is_new": save_as_new
        }
    except Exception as e:
        return {"success": False, "error": f"Failed to save document: {str(e)}"}

@app.route("/")
def index():
    return render_template("index.html")

@app.route("/api/config", methods=["GET"])
def get_config():
    cfg = config_manager.load_config()
    email_val = cfg.get("email") or cfg.get("gmail_address", "")
    cfg["email"] = email_val
    cfg["gmail_address"] = email_val
    cfg["telegram_bot_token"] = cfg.get("telegram_bot_token") or cfg.get("telegram_token", "")
    return jsonify(cfg)

@app.route("/api/config", methods=["POST"])
def save_config_endpoint():
    try:
        data = request.json or {}
        email_val = data.get("email") or data.get("gmail_address", "")
        data["email"] = email_val
        data["gmail_address"] = email_val
        if "telegram_bot_token" in data:
            data["telegram_token"] = data["telegram_bot_token"]
        config_manager.save_config(data)
        add_log("Configuration updated via recruitment hub.", "INFO")
        return jsonify({"success": True})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)})

@app.route("/api/startup/status", methods=["GET"])
def startup_status():
    return jsonify({"enabled": is_startup_enabled()})

@app.route("/api/startup/toggle", methods=["POST"])
def startup_toggle():
    try:
        data = request.json or {}
        enable = data.get("enable", False)
        if enable:
            ok = enable_startup()
            msg = "Silent Windows Startup enabled." if ok else "Failed to enable startup."
        else:
            ok = disable_startup()
            msg = "Silent Windows Startup disabled." if ok else "Failed to disable startup."
        return jsonify({"success": ok, "message": msg, "enabled": is_startup_enabled()})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)})

@app.route("/api/test_toast", methods=["POST"])
def test_toast_endpoint():
    send_windows_toast(
        title="📄 Resume Downloader Active",
        message="Windows Desktop Toast is working perfectly with sound!"
    )
    add_log("Dispatched test Windows Desktop Toast.", "INFO")
    return jsonify({"success": True})

@app.route("/api/test_telegram", methods=["POST"])
def test_telegram_endpoint():
    cfg = config_manager.load_config()
    tok = cfg.get("telegram_bot_token") or cfg.get("telegram_token", "")
    cid = cfg.get("telegram_chat_id", "")
    ok, msg = send_telegram_alert(tok, cid, "🚀 *Resume Downloader Alert Test*\n\nYour recruitment alert system is active!")
    add_log(f"Dispatched test Telegram alert. Result: {msg}", "INFO" if ok else "WARNING")
    return jsonify({"success": ok, "message": msg})

@app.route("/api/view_pdf", methods=["GET"])
def view_pdf_endpoint():
    file_path = request.args.get("path", "")
    if not file_path or not os.path.exists(file_path):
        return Response("Resume PDF file not found.", status=404)
    if not file_path.lower().endswith(".pdf"):
        return Response("Only PDF documents can be previewed directly.", status=400)
    return send_file(file_path, mimetype="application/pdf")

@app.route("/api/document/read", methods=["GET"])
def read_doc_endpoint():
    file_path = request.args.get("path", "")
    res = read_resume_document(file_path)
    return jsonify(res)

@app.route("/api/document/save", methods=["POST"])
def save_doc_endpoint():
    data = request.json or {}
    file_path = data.get("file_path", "")
    content = data.get("content", "")
    save_as_new = data.get("save_as_new", False)

    res = save_resume_document(file_path, content, save_as_new)
    if res.get("success"):
        cfg = config_manager.load_config()
        base_dir = cfg.get("base_directory", app_dir)
        fname = res.get("filename")
        add_log(f"Edited resume '{fname}' and saved changes to local disk.", "SUCCESS")
    return jsonify(res)

@app.route("/api/document/open_native", methods=["POST"])
def open_native_endpoint():
    data = request.json or {}
    file_path = data.get("file_path", "")
    if not file_path:
        return jsonify({"success": False, "error": "No file path provided."})
    ext = os.path.splitext(file_path)[1].lower()
    if ext not in SAFE_DOC_EXTENSIONS:
        return jsonify({"success": False, "error": "Only resume documents can be opened."})
    if os.path.exists(file_path):
        os.startfile(file_path)
        add_log(f"Opened '{os.path.basename(file_path)}' in native desktop application.", "INFO")
        return jsonify({"success": True})
    return jsonify({"success": False, "error": "File not found."})

@app.route("/api/test_connection", methods=["POST"])
def test_connection_endpoint():
    data = request.json or {}
    email_addr = (data.get("email") or "").strip()
    app_pwd = "".join((data.get("password") or data.get("app_password", "")).split())
    if not email_addr or not app_pwd:
        cfg = config_manager.load_config()
        email_addr = (cfg.get("email") or cfg.get("gmail_address", "")).strip()
        app_pwd = "".join((cfg.get("app_password", "")).split())
    if not email_addr or not app_pwd:
        return jsonify({"success": False, "error": "Email and 16-character App Password are required."})
    try:
        with MailBox("imap.gmail.com").login(email_addr, app_pwd):
            pass
        add_log("IMAP connection test successful for " + email_addr, "SUCCESS")
        return jsonify({"success": True})
    except Exception as e:
        add_log(f"IMAP connection test failed for {email_addr}: {e}", "ERROR")
        return jsonify({"success": False, "error": str(e)})

@app.route("/api/manifest", methods=["GET"])
def get_manifest():
    cfg = config_manager.load_config()
    base_dir = cfg.get("base_directory", app_dir)
    records = get_manifest_records(base_dir)
    return jsonify({"records": records})

@app.route("/api/manifest/update_status", methods=["POST"])
def update_status_endpoint():
    try:
        data = request.json or {}
        file_path = data.get("file_path", "")
        new_status = data.get("status", "New")
        notes = data.get("notes", "")

        cfg = config_manager.load_config()
        base_dir = cfg.get("base_directory", app_dir)
        ok = update_manifest_row(base_dir, file_path, new_status, notes)
        if ok:
            add_log(f"Updated candidate '{os.path.basename(file_path)}' status to '{new_status}'.", "INFO")
            return jsonify({"success": True})
        return jsonify({"success": False, "error": "Could not find matching candidate record."})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)})

@app.route("/api/manifest/export", methods=["GET"])
def export_manifest():
    cfg = config_manager.load_config()
    base_dir = cfg.get("base_directory", app_dir)
    manifest_path = os.path.join(base_dir, "resumes_manifest.csv")
    if os.path.exists(manifest_path):
        return send_file(manifest_path, as_attachment=True, download_name="resumes_manifest.csv")
    return Response("Downloaded At,Company,Applicant Name,Applicant Email,Email Subject,Filename,File Size (KB),File Path,SHA256,Status,Notes\n", mimetype="text/csv")

@app.route("/api/manifest/export_zip", methods=["GET"])
def export_zip_endpoint():
    try:
        company = request.args.get("company", "all").strip()
        cfg = config_manager.load_config()
        base_dir = cfg.get("base_directory", app_dir)

        records = get_manifest_records(base_dir)
        if company and company.lower() != "all":
            filtered_records = [r for r in records if r.get("Company", "").lower() == company.lower()]
            zip_filename = f"resumes_{company}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.zip"
        else:
            filtered_records = records
            zip_filename = f"all_resumes_{datetime.now().strftime('%Y%m%d_%H%M%S')}.zip"

        zip_buffer = io.BytesIO()
        with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
            csv_str_io = io.StringIO()
            writer = csv.DictWriter(csv_str_io, fieldnames=MANIFEST_HEADERS)
            writer.writeheader()
            for r in filtered_records:
                writer.writerow(r)
            zip_file.writestr("resumes_manifest_summary.csv", csv_str_io.getvalue())

            added_paths = set()
            for r in filtered_records:
                fpath = r.get("File Path", "")
                if fpath and os.path.exists(fpath) and fpath not in added_paths:
                    arcname = os.path.join(r.get("Company", "General"), r.get("Filename", os.path.basename(fpath)))
                    zip_file.write(fpath, arcname=arcname)
                    added_paths.add(fpath)

        zip_buffer.seek(0)
        return send_file(
            zip_buffer,
            mimetype="application/zip",
            as_attachment=True,
            download_name=zip_filename
        )
    except Exception as e:
        add_log(f"Error generating ZIP archive: {e}", "ERROR")
        return Response(f"Failed to generate ZIP: {e}", status=500)

@app.route("/api/scan_now", methods=["POST"])
def scan_now_endpoint():
    global active_watcher
    with watcher_lock:
        if active_watcher and active_watcher.thread and active_watcher.thread.is_alive():
            active_watcher.trigger_scan_now()
            add_log("Manual instant scan requested. Checking for new resumes...", "INFO")
            return jsonify({"success": True, "message": "Instant scan initiated!"})
        else:
            return jsonify({"success": False, "message": "Monitoring engine is stopped. Please start ingestion first."})

@app.route("/api/open_folder", methods=["POST"])
def open_base_folder():
    cfg = config_manager.load_config()
    base_dir = cfg.get("base_directory", app_dir)
    if os.path.exists(base_dir):
        os.startfile(base_dir)
        return jsonify({"success": True})
    return jsonify({"success": False, "error": "Base directory does not exist."})

@app.route("/api/open_file", methods=["POST"])
def open_file():
    data = request.json or {}
    fpath = data.get("path", "")
    if not fpath:
        return jsonify({"success": False, "error": "No file path provided."})
    ext = os.path.splitext(fpath)[1].lower()
    if ext not in SAFE_DOC_EXTENSIONS:
        return jsonify({"success": False, "error": "Only resume documents can be opened."})
    if os.path.exists(fpath):
        os.startfile(fpath)
        return jsonify({"success": True})
    return jsonify({"success": False, "error": "File not found."})

@app.route("/api/monitoring/status", methods=["GET"])
def monitoring_status():
    global active_watcher
    with watcher_lock:
        is_active = (active_watcher is not None and active_watcher.thread is not None and active_watcher.thread.is_alive())
        dups = active_watcher.duplicates_prevented if active_watcher else 0
        cls_name = active_watcher.__class__.__name__ if active_watcher else "None"
        mode = "cloud_sync" if cls_name == "FolderWatcher" else ("direct_gmail" if cls_name == "EmailWatcher" else "idle")
    with log_lock:
        recent_logs = list(log_messages[-100:])
    return jsonify({
        "is_monitoring": is_active,
        "mode": mode,
        "log_count": len(log_messages),
        "duplicates_prevented": dups,
        "logs": recent_logs,
        "stats": {
            "duplicates_prevented": dups
        }
    })

@app.route("/api/monitoring/start", methods=["POST"])
def start_monitoring():
    global active_watcher, email_watcher
    req_data = request.json or {}
    cfg = config_manager.load_config()

    mode = req_data.get("mode") or cfg.get("monitoring_mode", "cloud_sync")
    base_dir = (req_data.get("base_directory") or cfg.get("base_directory", app_dir)).strip()
    
    with watcher_lock:
        if active_watcher and active_watcher.thread and active_watcher.thread.is_alive():
            return jsonify({"success": True, "message": "Monitoring is already active."})

        if mode == "cloud_sync":
            if not base_dir:
                return jsonify({"success": False, "error": "Please configure Base / Google Drive Synced Directory."})
            os.makedirs(base_dir, exist_ok=True)
            active_watcher = FolderWatcher(
                watch_dir=base_dir,
                companies=cfg.get("companies", []),
                create_subfolders=cfg.get("create_subfolders", True),
                enable_desktop_toasts=cfg.get("enable_desktop_toasts", True),
                telegram_token=cfg.get("telegram_bot_token") or cfg.get("telegram_token", ""),
                telegram_chat_id=cfg.get("telegram_chat_id", ""),
                log_callback=add_log,
                poll_interval=5
            )
            email_watcher = active_watcher
            active_watcher.start()
            add_log(f"☁️ Cloud Sync Folder Watcher started. Monitoring: '{base_dir}' with Auto Title-Sorting.", "SUCCESS")
            return jsonify({"success": True, "mode": "cloud_sync"})

        else:
            # direct_gmail mode (supports multiple accounts)
            accounts = cfg.get("accounts", [])
            if not accounts:
                email_addr = (cfg.get("email") or cfg.get("gmail_address", "")).strip()
                app_pwd = "".join((cfg.get("app_password", "")).split())
                if email_addr and app_pwd:
                    accounts = [{"email": email_addr, "app_password": app_pwd}]

            valid_accounts = [
                a for a in accounts
                if isinstance(a, dict) and a.get("email") and a.get("app_password")
            ]

            if not valid_accounts or not base_dir:
                return jsonify({
                    "success": False,
                    "error": "Direct Gmail mode requires at least one configured Gmail account (with 16-char App Password) and Local Directory."
                })

            active_watcher = EmailWatcher(
                accounts=valid_accounts,
                base_dir=base_dir,
                companies=cfg.get("companies", []),
                create_subfolders=cfg.get("create_subfolders", True),
                prefix_applicant_name=cfg.get("prefix_applicant_name", True),
                enable_desktop_toasts=cfg.get("enable_desktop_toasts", True),
                telegram_token=cfg.get("telegram_bot_token") or cfg.get("telegram_token", ""),
                telegram_chat_id=cfg.get("telegram_chat_id", ""),
                log_callback=add_log,
                poll_interval=10
            )
            email_watcher = active_watcher
            active_watcher.start()
            acc_list_str = ", ".join([a.get("email", "") for a in valid_accounts])
            add_log(f"✉️ Autonomous Gmail IMAP monitoring started for {len(valid_accounts)} account(s): [{acc_list_str}].", "SUCCESS")
            return jsonify({"success": True, "mode": "direct_gmail", "accounts_count": len(valid_accounts)})

def stop_monitoring_internal():
    global active_watcher, email_watcher
    with watcher_lock:
        if active_watcher:
            active_watcher.stop()
            active_watcher = None
            email_watcher = None

@app.route("/api/monitoring/stop", methods=["POST"])
def stop_monitoring():
    stop_monitoring_internal()
    add_log("Monitoring stopped.", "INFO")
    return jsonify({"success": True})

@app.route("/api/logs/clear", methods=["POST"])
def clear_logs():
    with log_lock:
        log_messages.clear()
    return jsonify({"success": True})

@app.route("/api/logs/stream")
def stream_logs():
    def event_stream():
        last_index = 0
        while True:
            with log_lock:
                current_len = len(log_messages)
                if last_index < current_len:
                    for i in range(last_index, current_len):
                        yield f"data: {json.dumps(log_messages[i])}\n\n"
                    last_index = current_len
            time.sleep(1)
    return Response(event_stream(), mimetype="text/event-stream")

if __name__ == "__main__":
    add_log("Enterprise Recruitment Hub initialized on http://127.0.0.1:5000", "SUCCESS")
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port, debug=False, threaded=True)
